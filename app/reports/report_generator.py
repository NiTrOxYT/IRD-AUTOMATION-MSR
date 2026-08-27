import os
import csv
import json
import logging
from datetime import datetime
from typing import Dict, Any, List
from pathlib import Path

from app.config import REPORTS_DIR

logger = logging.getLogger("IRD_ReportGenerator")

def generate_text_summary(run_data: Dict[str, Any]) -> str:
    """Generates formatted ASCII report as specified in Section 19."""
    run = run_data.get("run", {})
    results = run_data.get("results", [])

    date_str = run.get("run_date", datetime.now().strftime("%Y-%m-%d"))
    total = run.get("total_sites", len(results))
    completed = run.get("completed_sites", 0)

    no_action_sites = [r["site_name"] for r in results if r.get("action_point_status") == "NO ACTION POINT"]
    action_point_sites = [r["site_name"] for r in results if r.get("action_point_status") == "ACTION POINT FOUND"]
    failed_sites = [r["site_name"] for r in results if r.get("action_point_status") == "FAILED" or r.get("status") == "FAILED"]

    report_lines = [
        "======================================",
        "          IRD SYNC REPORT             ",
        f"Date: {date_str}",
        "======================================",
        f"Total Sites: {total}",
        f"Completed:   {completed}",
        "--------------------------------------",
        "",
        "NO ACTION POINT",
        ""
    ]

    if no_action_sites:
        for s in no_action_sites:
            report_lines.append(f"• {s}")
    else:
        report_lines.append("None")

    report_lines.extend([
        "",
        "--------------------------------------",
        "",
        "ACTION POINT FOUND",
        ""
    ])

    if action_point_sites:
        for s in action_point_sites:
            report_lines.append(f"• {s}")
    else:
        report_lines.append("None")

    report_lines.extend([
        "",
        "--------------------------------------",
        "",
        "FAILED",
        ""
    ])

    if failed_sites:
        for s in failed_sites:
            report_lines.append(f"• {s}")
    else:
        report_lines.append("None")

    report_lines.extend([
        "",
        "--------------------------------------",
        f"Action Point Found: {len(action_point_sites)}",
        f"No Action Point:    {len(no_action_sites)}",
        f"Failed:             {len(failed_sites)}",
        "======================================"
    ])

    return "\n".join(report_lines)


def export_reports(run_data: Dict[str, Any], output_dir: str = None) -> Dict[str, str]:
    """
    Exports run results to TXT, CSV, Excel (.xlsx), and JSON.
    Returns dict of filepaths.
    """
    out_folder = Path(output_dir) if output_dir else REPORTS_DIR
    out_folder.mkdir(parents=True, exist_ok=True)

    run = run_data.get("run", {})
    results = run_data.get("results", [])
    date_str = run.get("run_date", datetime.now().strftime("%Y-%m-%d"))

    txt_path = out_folder / "daily_report.txt"
    csv_path = out_folder / "daily_report.csv"
    xlsx_path = out_folder / "daily_report.xlsx"
    json_path = out_folder / "daily_report.json"


    # 1. Write Text Report
    text_summary = generate_text_summary(run_data)
    with open(txt_path, "w", encoding="utf-8") as f:
        f.write(text_summary)

    # 2. Write CSV Report
    fieldnames = [
        "Date", "Site", "Status", "Action Point Status", "CSV File",
        "Third Line", "Start Time", "End Time", "Duration (s)", "Retries", "Error"
    ]

    rows = []
    for r in results:
        rows.append({
            "Date": date_str,
            "Site": r.get("site_name", ""),
            "Status": r.get("status", ""),
            "Action Point Status": r.get("action_point_status", ""),
            "CSV File": r.get("csv_path", ""),
            "Third Line": r.get("third_line_text", ""),
            "Start Time": r.get("start_time", ""),
            "End Time": r.get("end_time", ""),
            "Duration (s)": round(r.get("duration_seconds", 0.0), 2),
            "Retries": r.get("retry_count", 0),
            "Error": r.get("error_message", "")
        })

    with open(csv_path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)

    # 3. Write JSON Report
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(run_data, f, indent=2)

    # 4. Write Excel (.xlsx) Report using openpyxl
    try:
        from openpyxl import Workbook
        from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

        wb = Workbook()
        ws = wb.active
        ws.title = "IRD Sync Report"

        # Headers
        ws.append(fieldnames)

        # Styles
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

        green_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid") # No Action
        amber_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid") # Action Point Found
        red_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")   # Failed

        for col_num in range(1, len(fieldnames) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for r_idx, row_dict in enumerate(rows, start=2):
            ws.append([row_dict[k] for k in fieldnames])
            ap_stat = row_dict["Action Point Status"]
            fill_to_use = green_fill if ap_stat == "NO ACTION POINT" else (amber_fill if ap_stat == "ACTION POINT FOUND" else red_fill)
            for c_idx in range(1, len(fieldnames) + 1):
                ws.cell(row=r_idx, column=c_idx).fill = fill_to_use

        # Auto column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = cell.column_letter
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        wb.save(xlsx_path)
    except Exception as ex:
        logger.error(f"Excel report generation failed: {ex}")

    logger.info(f"Generated reports in folder: {out_folder}")

    return {
        "txt": str(txt_path),
        "csv": str(csv_path),
        "xlsx": str(xlsx_path),
        "json": str(json_path)
    }
