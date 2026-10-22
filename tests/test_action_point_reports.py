import os
import json
import csv
import pytest
import openpyxl

from app.action_point_automation.models import (
    ItemExecutionResult,
    ActionClassification,
    ActionPointStatusEnum,
    ExecutionMode
)
from app.action_point_automation.report_generator import export_action_point_reports


def test_export_action_point_reports():
    results = [
        ItemExecutionResult(
            row_number=2,
            item_code="0042",
            item_name="Chicken Tikka",
            requested_category="Starters",
            requested_subcategory="Non-Veg Starters",
            original_status="Category Update",
            classification=ActionClassification.CATEGORY_UPDATE,
            status=ActionPointStatusEnum.SUCCESS,
            category_created=False,
            subcategory_verified=True,
            item_found=True,
            restaurant_verified=True,
            portal_verified=True,
            failure_reason=""
        ),
        ItemExecutionResult(
            row_number=3,
            item_code="0043",
            item_name="Cold Coffee",
            requested_category="Beverages",
            requested_subcategory="Cold Drinks",
            original_status="Active",
            classification=ActionClassification.UNSUPPORTED_STATUS,
            status=ActionPointStatusEnum.SKIPPED_UNSUPPORTED_STATUS,
            failure_reason="Unsupported status 'Active'"
        )
    ]

    report_paths = export_action_point_reports(
        run_id=9999,
        site_name="Test_Maurya",
        filename="test_menu.xlsx",
        mode=ExecutionMode.LIVE,
        results=results,
        summary_stats={"total_rows": 2, "successful_rows": 1, "skipped_rows": 1}
    )

    csv_path = report_paths["csv"]
    json_path = report_paths["json"]
    xlsx_path = report_paths["xlsx"]

    assert os.path.exists(csv_path)
    assert os.path.exists(json_path)
    assert os.path.exists(xlsx_path)

    # 1. Verify JSON
    with open(json_path, "r", encoding="utf-8") as f:
        data = json.load(f)
    assert data["run_id"] == 9999
    assert data["site_name"] == "Test_Maurya"
    assert len(data["items"]) == 2

    # 2. Verify CSV
    with open(csv_path, "r", encoding="utf-8-sig") as f:
        reader = list(csv.reader(f))
    assert len(reader) == 3  # Header + 2 rows
    assert reader[1][1] == "0042"
    assert reader[1][7] == "SUCCESS"
    assert reader[1][9] == "NO"  # Subcategory Created
    assert reader[2][7] == "SKIPPED_UNSUPPORTED_STATUS"

    # 3. Verify XLSX
    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb.active
    assert ws.cell(row=2, column=2).value == "0042"
    assert ws.cell(row=2, column=8).value == "SUCCESS"
    assert ws.cell(row=2, column=10).value == "NO"  # Subcategory Created
    assert ws.cell(row=3, column=8).value == "SKIPPED_UNSUPPORTED_STATUS"
    wb.close()

    # Clean up test files
    for p in [csv_path, json_path, xlsx_path]:
        try:
            if os.path.exists(p):
                os.remove(p)
        except Exception:
            pass
