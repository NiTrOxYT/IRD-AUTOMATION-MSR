import os
import csv
import json
import logging
from datetime import datetime
from pathlib import Path
from typing import List, Dict, Any, Optional

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from app.config import REPORTS_DIR
from app.action_point_automation.models import ItemExecutionResult, ExecutionMode

logger = logging.getLogger("ActionPoint_ReportGenerator")


def export_action_point_reports(
    run_id: int,
    site_name: str,
    filename: str,
    mode: ExecutionMode,
    results: List[ItemExecutionResult],
    summary_stats: Optional[Dict[str, Any]] = None
) -> Dict[str, str]:
    """
    Exports action point configuration results in CSV, JSON, and XLSX formats.
    Returns: Dict mapping format ('csv', 'json', 'xlsx') to generated file path.
    """
    date_str = datetime.now().strftime("%Y-%m-%d")
    clean_site = "".join(c if c.isalnum() or c in ("-", "_") else "_" for c in site_name)
    target_dir = REPORTS_DIR / "action_point"
    target_dir.mkdir(parents=True, exist_ok=True)

    base_name = f"{date_str}_{clean_site}_run_{run_id}_{mode.value.lower()}"
    csv_path = target_dir / f"{base_name}.csv"
    json_path = target_dir / f"{base_name}.json"
    xlsx_path = target_dir / f"{base_name}.xlsx"

    # 1. Export JSON Report
    report_data = {
        "run_id": run_id,
        "site_name": site_name,
        "source_filename": filename,
        "execution_mode": mode.value,
        "generated_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "summary": summary_stats or {},
        "items": [res.model_dump() for res in results]
    }
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(report_data, f, indent=2, default=str)

    # 2. Export CSV Report
    headers = [
        "Row Number",
        "Item Code",
        "Item Name",
        "Requested Category",
        "Requested Subcategory",
        "Original Status",
        "Classification",
        "Result Status",
        "Item Found",
        "Item Checkbox Selected",
        "Category Selected",
        "Subcategory Selected",
        "Subcategory Skipped",
        "Restaurant Selected",
        "Configure Submitted",
        "Persisted Verified",
        "Category Created",
        "Subcategory Created",
        "Subcategory Verified",
        "Restaurant Verified",
        "Failure Stage",
        "Failure / Skip Reason",
        "Duration (s)",
        "Search Duration (s)",
        "Category Select Duration (s)",
        "Save Duration (s)",
        "Post-Save Duration (s)",
        "Screenshot Path"
    ]
    with open(csv_path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(headers)
        for r in results:
            writer.writerow([
                r.row_number,
                r.item_code,
                r.item_name,
                r.requested_category,
                r.requested_subcategory,
                r.original_status,
                r.classification.value if hasattr(r.classification, "value") else str(r.classification),
                r.status.value if hasattr(r.status, "value") else str(r.status),
                "YES" if getattr(r, "item_found", False) else "NO",
                "YES" if getattr(r, "item_checkbox_selected", False) else "NO",
                "YES" if getattr(r, "category_selected", False) else "NO",
                "YES" if getattr(r, "subcategory_selected", False) else "NO",
                "YES" if getattr(r, "subcategory_skipped", False) else "NO",
                "YES" if getattr(r, "restaurant_selected", False) else "NO",
                "YES" if getattr(r, "configure_submitted", False) else "NO",
                "YES" if getattr(r, "persisted_verified", False) else "NO",
                "YES" if r.category_created else "NO",
                "YES" if getattr(r, "subcategory_created", False) else "NO",
                "YES" if r.subcategory_verified else "NO",
                "YES" if r.restaurant_verified else "NO",
                r.failure_stage or "",
                r.failure_reason,
                r.duration_seconds,
                getattr(r, "search_duration_seconds", 0.0),
                getattr(r, "category_select_duration_seconds", 0.0),
                getattr(r, "save_duration_seconds", 0.0),
                getattr(r, "post_save_duration_seconds", 0.0),
                r.screenshot_path or ""
            ])

    # 3. Export XLSX Report with professional styling
    try:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Action Point Execution"

        # Header styling
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        thin_border = Border(
            left=Side(style="thin", color="CBD5E1"),
            right=Side(style="thin", color="CBD5E1"),
            top=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="thin", color="CBD5E1")
        )

        ws.append(headers)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

        for r in results:
            row_data = [
                r.row_number,
                r.item_code,
                r.item_name,
                r.requested_category,
                r.requested_subcategory,
                r.original_status,
                r.classification.value if hasattr(r.classification, "value") else str(r.classification),
                r.status.value if hasattr(r.status, "value") else str(r.status),
                "YES" if getattr(r, "item_found", False) else "NO",
                "YES" if getattr(r, "item_checkbox_selected", False) else "NO",
                "YES" if getattr(r, "category_selected", False) else "NO",
                "YES" if getattr(r, "subcategory_selected", False) else "NO",
                "YES" if getattr(r, "subcategory_skipped", False) else "NO",
                "YES" if getattr(r, "restaurant_selected", False) else "NO",
                "YES" if getattr(r, "configure_submitted", False) else "NO",
                "YES" if getattr(r, "persisted_verified", False) else "NO",
                "YES" if r.category_created else "NO",
                "YES" if getattr(r, "subcategory_created", False) else "NO",
                "YES" if r.subcategory_verified else "NO",
                "YES" if r.restaurant_verified else "NO",
                r.failure_stage or "",
                r.failure_reason,
                r.duration_seconds,
                getattr(r, "search_duration_seconds", 0.0),
                getattr(r, "category_select_duration_seconds", 0.0),
                getattr(r, "save_duration_seconds", 0.0),
                getattr(r, "post_save_duration_seconds", 0.0),
                r.screenshot_path or ""
            ]
            ws.append(row_data)

            # Color-code result status cell
            row_num = ws.max_row
            status_cell = ws.cell(row=row_num, column=8)
            status_val = status_cell.value

            if status_val in ("SUCCESS", "DRY_RUN"):
                status_cell.fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
                status_cell.font = Font(name="Calibri", bold=True, color="166534")
            elif "SKIPPED" in str(status_val):
                status_cell.fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
                status_cell.font = Font(name="Calibri", bold=False, color="475569")
            elif status_val == "BLOCKED":
                status_cell.fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
                status_cell.font = Font(name="Calibri", bold=True, color="92400E")
            elif status_val in ("FAILED", "VERIFICATION_FAILED"):
                status_cell.fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
                status_cell.font = Font(name="Calibri", bold=True, color="991B1B")

            for col_idx in range(1, len(headers) + 1):
                ws.cell(row=row_num, column=col_idx).border = thin_border

        # Adjust column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        # Add Performance Metrics worksheet if present
        perf = (summary_stats or {}).get("performance_metrics")
        if perf:
            ws_perf = wb.create_sheet(title="Performance Metrics")
            perf_headers = ["Metric Description", "Measured Value"]
            ws_perf.append(perf_headers)
            for col_idx in range(1, 3):
                c = ws_perf.cell(row=1, column=col_idx)
                c.fill = header_fill
                c.font = header_font
                c.alignment = Alignment(horizontal="center", vertical="center")
                c.border = thin_border

            metric_rows = [
                ("Total Run Duration (seconds)", f"{perf.get('total_run_duration_seconds', 0.0):.2f}"),
                ("Average Time Per Item (seconds)", f"{perf.get('average_time_per_item_seconds', 0.0):.2f}"),
                ("Portal Navigation Duration (seconds)", f"{perf.get('portal_navigation_duration_seconds', 0.0):.2f}"),
                ("Login and Auth Duration (seconds)", f"{perf.get('login_and_auth_duration_seconds', 0.0):.2f}"),
                ("Category Inspection Duration (seconds)", f"{perf.get('category_inspection_duration_seconds', 0.0):.2f}"),
                ("Category Creation Duration (seconds)", f"{perf.get('category_creation_duration_seconds', 0.0):.2f}"),
                ("Subcategory Expansion Duration (seconds)", f"{perf.get('subcategory_expansion_duration_seconds', 0.0):.2f}"),
                ("Subcategory Creation Duration (seconds)", f"{perf.get('subcategory_creation_duration_seconds', 0.0):.2f}"),
                ("Item Processing Duration (seconds)", f"{perf.get('item_processing_duration_seconds', 0.0):.2f}"),
                ("Item Search Total (seconds)", f"{perf.get('item_search_total_seconds', 0.0):.2f}"),
                ("Category Select Total (seconds)", f"{perf.get('category_select_total_seconds', 0.0):.2f}"),
                ("Restaurant Select Total (seconds)", f"{perf.get('restaurant_select_duration_total', 0.0):.2f}"),
                ("Configure Save Total (seconds)", f"{perf.get('configure_save_total_seconds', 0.0):.2f}"),
                ("Post-Save Verification Total (seconds)", f"{perf.get('post_save_verification_total_seconds', 0.0):.2f}"),
                ("Screenshot Capture Total (seconds)", f"{perf.get('screenshot_capture_total_seconds', 0.0):.2f}"),
                ("Total Page Navigations", str(perf.get('total_navigations', 0))),
                ("Total Item Searches", str(perf.get('total_searches', 0))),
                ("Total Retries / Recoveries", str(perf.get('total_retries', 0))),
                ("Total Diagnostic Screenshots", str(perf.get('total_screenshots', 0)))
            ]
            for label, val in metric_rows:
                ws_perf.append([label, val])
                r_num = ws_perf.max_row
                ws_perf.cell(row=r_num, column=1).border = thin_border
                ws_perf.cell(row=r_num, column=2).border = thin_border

            for col in ws_perf.columns:
                max_len = max(len(str(cell.value or "")) for cell in col)
                col_letter = openpyxl.utils.get_column_letter(col[0].column)
                ws_perf.column_dimensions[col_letter].width = max(max_len + 3, 24)

        wb.save(str(xlsx_path))
    except Exception as e:
        logger.warning(f"Excel report generation encountered error: {e}")

    logger.info(f"[{site_name}] Action point reports generated: CSV={csv_path}, JSON={json_path}, XLSX={xlsx_path}")
    return {
        "csv": str(csv_path),
        "json": str(json_path),
        "xlsx": str(xlsx_path)
    }
